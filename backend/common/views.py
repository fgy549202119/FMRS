"""
公共模块 - 视图函数定义

该模块提供了公共功能相关的API视图，包括：
- 统计数据（StatsViewSet）
- 备件管理（SparePartViewSet）
- 备件记录（SparePartRecordViewSet）
- 维修日志（RepairLogViewSet）
- 维修备件使用（RepairSparePartViewSet）
- 知识库（KnowledgeBaseViewSet）
- 数据导出（ExportViewSet）

视图集结构：
    StatsViewSet
        - 系统统计数据
    
    SparePartViewSet
        - 备件CRUD操作
        - 入库、出库功能
        - 库存预警
    
    SparePartRecordViewSet
        - 备件记录查询
    
    RepairLogViewSet
        - 维修日志管理
    
    RepairSparePartViewSet
        - 维修备件使用管理
    
    KnowledgeBaseViewSet
        - 知识库CRUD操作
        - 浏览计数、有用计数
    
    ExportViewSet
        - Excel导出
        - PDF导出

作者：FMRS开发者范广宇
创建日期：2026年
"""

from rest_framework import viewsets, status
from rest_framework.decorators import action
from rest_framework.response import Response
from django_filters.rest_framework import DjangoFilterBackend
from rest_framework import filters
from django.utils import timezone
from django.http import HttpResponse
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side
from fpdf import FPDF
import io
import os

from .models import SparePart, SparePartRecord, RepairLog, RepairSparePart, KnowledgeBase
from .serializers import (
    SparePartSerializer, SparePartRecordSerializer,
    RepairLogSerializer, RepairSparePartSerializer, KnowledgeBaseSerializer
)
from repair.models import RepairOrder
from users.models import User, RepairStaff
from equipment.models import Equipment


class StatsViewSet(viewsets.ViewSet):
    """
    统计数据视图集
    
    提供系统级别的统计数据。
    
    自定义操作：
        - summary: 系统概览统计
    """
    
    @action(detail=False, methods=['get'])
    def summary(self, request):
        """
        系统概览统计接口
        
        返回系统各项数据的统计信息。
        
        请求方式：
            GET /api/common/stats/summary/
        
        返回格式：
            {
                'code': 200,
                'data': {
                    'user_count': 用户总数,
                    'equipment_count': 设备总数,
                    'repair_order_count': 工单总数,
                    'staff_count': 维修人员总数,
                    'repair_status': {各状态工单数量}
                }
            }
        """
        from repair.models import RepairOrder
        data = {
            'user_count': User.objects.count(),
            'equipment_count': Equipment.objects.count(),
            'repair_order_count': RepairOrder.objects.filter(is_deleted=False).count(),
            'staff_count': RepairStaff.objects.count(),
            'repair_status': {
                'pending': RepairOrder.objects.filter(status='WAITING', is_deleted=False).count(),
                'accepted': RepairOrder.objects.filter(status='IN_PROGRESS', is_deleted=False).count(),
                'repairing': RepairOrder.objects.filter(status='PENDING_ADMIN_CLOSE', is_deleted=False).count(),
                'completed': RepairOrder.objects.filter(status='CLOSED', is_deleted=False).count(),
                'rejected': RepairOrder.objects.filter(status='REJECTED', is_deleted=False).count(),
            }
        }
        return Response({'code': 200, 'data': data})


class SparePartViewSet(viewsets.ModelViewSet):
    """
    备件视图集
    
    提供备件的完整CRUD操作，以及入库、出库功能。
    
    继承自ModelViewSet，自动提供CRUD操作。
    
    过滤功能：
        - 按备件编号搜索
        - 按备件名称搜索
        - 按规格型号搜索
    
    自定义操作：
        - low_stock: 获取库存不足的备件
        - stock_in: 备件入库
        - stock_out: 备件出库
    """
    
    # 查询集
    queryset = SparePart.objects.all()
    
    # 序列化器
    serializer_class = SparePartSerializer
    
    # 过滤后端
    filter_backends = [DjangoFilterBackend, filters.SearchFilter]
    
    # 可搜索字段
    search_fields = ['part_no', 'name', 'specification']

    @action(detail=False, methods=['get'])
    def low_stock(self, request):
        """
        获取库存不足的备件列表
        
        返回库存数量<=0的备件列表。
        
        请求方式：
            GET /api/common/spare-parts/low_stock/
        
        返回格式：
            [{备件信息}]
        """
        parts = SparePart.objects.filter(quantity__lte=0)
        serializer = self.get_serializer(parts, many=True)
        return Response(serializer.data)
    
    @action(detail=True, methods=['post'])
    def stock_in(self, request, pk=None):
        """
        备件入库接口
        
        增加备件库存数量，并记录入库历史。
        
        请求方式：
            POST /api/common/spare-parts/{id}/stock_in/
        
        请求参数：
            - quantity: 入库数量
            - remark: 备注
            - operator: 操作人
        
        返回格式：
            {'code': 200, 'message': '入库成功', 'quantity': 当前库存}
        """
        part = self.get_object()
        quantity = request.data.get('quantity', 0)
        remark = request.data.get('remark', '')
        operator = request.data.get('operator', '')
        
        # 记录入库前库存
        before = part.quantity
        
        # 更新库存
        part.quantity += quantity
        part.save()
        
        # 创建入库记录
        SparePartRecord.objects.create(
            part=part,
            record_type='in',
            quantity=quantity,
            before_quantity=before,
            after_quantity=part.quantity,
            operator=operator,
            remark=remark
        )
        
        return Response({'code': 200, 'message': '入库成功', 'quantity': part.quantity})
    
    @action(detail=True, methods=['post'])
    def stock_out(self, request, pk=None):
        """
        备件出库接口
        
        减少备件库存数量，并记录出库历史。
        库存不足时会返回错误。
        
        请求方式：
            POST /api/common/spare-parts/{id}/stock_out/
        
        请求参数：
            - quantity: 出库数量
            - remark: 备注
            - operator: 操作人
            - related_order_no: 关联工单号
        
        返回格式：
            成功：{'code': 200, 'message': '出库成功', 'quantity': 当前库存}
            失败：{'code': 400, 'message': '库存不足'}
        """
        part = self.get_object()
        quantity = request.data.get('quantity', 0)
        remark = request.data.get('remark', '')
        operator = request.data.get('operator', '')
        related_order_no = request.data.get('related_order_no', '')
        
        # 检查库存是否充足
        if part.quantity < quantity:
            return Response({'code': 400, 'message': '库存不足'}, status=status.HTTP_400_BAD_REQUEST)
        
        # 记录出库前库存
        before = part.quantity
        
        # 更新库存
        part.quantity -= quantity
        part.save()
        
        # 创建出库记录
        SparePartRecord.objects.create(
            part=part,
            record_type='out',
            quantity=quantity,
            before_quantity=before,
            after_quantity=part.quantity,
            operator=operator,
            remark=remark,
            related_order_no=related_order_no
        )
        
        return Response({'code': 200, 'message': '出库成功', 'quantity': part.quantity})


class SparePartRecordViewSet(viewsets.ReadOnlyModelViewSet):
    """
    备件记录视图集
    
    提供备件记录的只读查询功能。
    
    继承自ReadOnlyModelViewSet，只提供查询操作。
    
    过滤功能：
        - 按备件过滤
        - 按记录类型过滤
        - 按关联单号搜索
        - 按操作人搜索
    """
    
    # 查询集
    queryset = SparePartRecord.objects.all()
    
    # 序列化器
    serializer_class = SparePartRecordSerializer
    
    # 过滤后端
    filter_backends = [DjangoFilterBackend, filters.SearchFilter]
    
    # 可过滤字段
    filterset_fields = ['part', 'record_type']
    
    # 可搜索字段
    search_fields = ['related_order_no', 'operator']


class RepairLogViewSet(viewsets.ModelViewSet):
    """
    维修日志视图集
    
    提供维修日志的完整CRUD操作。
    
    继承自ModelViewSet，自动提供CRUD操作。
    
    过滤功能：
        - 按工单过滤
    """
    
    # 查询集
    queryset = RepairLog.objects.all()
    
    # 序列化器
    serializer_class = RepairLogSerializer
    
    # 过滤后端
    filter_backends = [DjangoFilterBackend]
    
    # 可过滤字段
    filterset_fields = ['repair_order']


class RepairSparePartViewSet(viewsets.ModelViewSet):
    """
    维修备件使用视图集
    
    提供维修备件使用记录的完整CRUD操作。
    
    继承自ModelViewSet，自动提供CRUD操作。
    
    过滤功能：
        - 按工单过滤
        - 按备件过滤
    """
    
    # 查询集
    queryset = RepairSparePart.objects.all()
    
    # 序列化器
    serializer_class = RepairSparePartSerializer
    
    # 过滤后端
    filter_backends = [DjangoFilterBackend]
    
    # 可过滤字段
    filterset_fields = ['repair_order', 'spare_part']


class KnowledgeBaseViewSet(viewsets.ModelViewSet):
    """
    知识库视图集
    
    提供知识库的完整CRUD操作，以及浏览计数功能。
    
    继承自ModelViewSet，自动提供CRUD操作。
    
    过滤功能：
        - 按分类过滤
        - 按标题搜索
        - 按问题描述搜索
        - 按关键词搜索
    
    自定义操作：
        - view: 增加浏览次数
        - useful: 增加有用次数
    """
    
    # 查询集：只返回已发布的条目
    queryset = KnowledgeBase.objects.filter(is_published=True)
    
    # 序列化器
    serializer_class = KnowledgeBaseSerializer
    
    # 过滤后端
    filter_backends = [DjangoFilterBackend, filters.SearchFilter]
    
    # 可过滤字段
    filterset_fields = ['equipment_type']
    
    # 可搜索字段
    search_fields = ['title', 'problem_description', 'keywords']
    
    @action(detail=True, methods=['post'])
    def view(self, request, pk=None):
        """
        增加浏览次数
        
        请求方式：
            POST /api/common/knowledge/{id}/view/
        
        返回格式：
            {'code': 200, 'view_count': 当前浏览次数}
        """
        article = self.get_object()
        article.view_count += 1
        article.save()
        return Response({'code': 200, 'view_count': article.view_count})
    
    @action(detail=True, methods=['post'])
    def useful(self, request, pk=None):
        """
        增加有用次数
        
        请求方式：
            POST /api/common/knowledge/{id}/useful/
        
        返回格式：
            {'code': 200, 'useful_count': 当前有用次数}
        """
        article = self.get_object()
        article.useful_count += 1
        article.save()
        return Response({'code': 200, 'useful_count': article.useful_count})


class ExportViewSet(viewsets.ViewSet):
    """
    数据导出视图集
    
    提供数据导出功能，支持Excel和PDF格式。
    
    自定义操作：
        - repair_orders_excel: 导出工单Excel
        - repair_orders_pdf: 导出工单PDF
        - spare_parts_excel: 导出备件Excel
    """
    
    @action(detail=False, methods=['get'])
    def repair_orders_excel(self, request):
        """
        导出报修工单Excel
        
        支持按状态、日期范围筛选。
        
        请求方式：
            GET /api/common/export/repair_orders_excel/
        
        查询参数：
            - status: 工单状态
            - start_date: 开始日期
            - end_date: 结束日期
        
        返回：
            Excel文件下载
        """
        queryset = RepairOrder.objects.filter(is_deleted=False)
        
        # 获取筛选参数
        status_filter = request.query_params.get('status')
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')
        
        # 应用筛选
        if status_filter:
            queryset = queryset.filter(status=status_filter)
        if start_date:
            queryset = queryset.filter(created_at__gte=start_date)
        if end_date:
            queryset = queryset.filter(created_at__lte=end_date)
        
        # 创建Excel工作簿
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '报修工单'
        
        # 设置表头
        headers = ['报修编号', '设备名称', '设备类型', '位置', '报修人', '状态', '维修员', '创建时间', '完成时间', '耗时(分钟)']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal='center')
        
        # 状态映射
        status_map = {
            'WAITING': '待接单', 'IN_PROGRESS': '维修中',
            'PENDING_ADMIN_CLOSE': '待审核', 'CLOSED': '已完成', 'REJECTED': '已取消'
        }
        
        # 填充数据
        for row, order in enumerate(queryset, 2):
            ws.cell(row=row, column=1, value=order.repair_no)
            ws.cell(row=row, column=2, value=order.equipment_name)
            ws.cell(row=row, column=3, value=str(order.equipment_type) if order.equipment_type else '')
            ws.cell(row=row, column=4, value=f"{order.location or ''}{order.building or ''}{order.floor or ''}{order.room or ''}")
            ws.cell(row=row, column=5, value=order.user_name or '')
            ws.cell(row=row, column=6, value=status_map.get(order.status, order.status))
            ws.cell(row=row, column=7, value=order.staff_name or '')
            ws.cell(row=row, column=8, value=order.created_at.strftime('%Y-%m-%d %H:%M') if order.created_at else '')
            ws.cell(row=row, column=9, value=order.complete_time.strftime('%Y-%m-%d %H:%M') if order.complete_time else '')
            
            # 计算耗时
            if order.repair_start_time and order.complete_time:
                duration = (order.complete_time - order.repair_start_time).total_seconds() / 60
                ws.cell(row=row, column=10, value=round(duration, 1))
            else:
                ws.cell(row=row, column=10, value='')
        
        # 自动调整列宽
        for col in ws.columns:
            max_length = 0
            column = col[0].column_letter
            for cell in col:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            ws.column_dimensions[column].width = min(max_length + 2, 50)
        
        # 保存到内存缓冲区
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        # 返回Excel文件
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="repair_orders.xlsx"'
        return response
    
    @action(detail=False, methods=['get'])
    def repair_orders_pdf(self, request):
        """
        导出报修工单PDF
        
        支持按状态、日期范围筛选。
        
        请求方式：
            GET /api/common/export/repair_orders_pdf/
        
        查询参数：
            - status: 工单状态
            - start_date: 开始日期
            - end_date: 结束日期
        
        返回：
            PDF文件下载
        """
        queryset = RepairOrder.objects.filter(is_deleted=False)
        
        # 获取筛选参数
        status_filter = request.query_params.get('status')
        start_date = request.query_params.get('start_date')
        end_date = request.query_params.get('end_date')
        
        # 应用筛选
        if status_filter:
            queryset = queryset.filter(status=status_filter)
        if start_date:
            queryset = queryset.filter(created_at__gte=start_date)
        if end_date:
            queryset = queryset.filter(created_at__lte=end_date)
        
        # 注册中文字体
        font_path = os.path.join(os.environ.get('SystemRoot', r'C:\Windows'), 'Fonts', 'simhei.ttf')
        
        # 创建PDF文档（横向A4）
        pdf = FPDF(orientation='L', unit='mm', format='A4')
        pdf.set_auto_page_break(auto=True, margin=10)
        
        # 添加中文字体支持
        if os.path.exists(font_path):
            pdf.add_font('SimHei', '', font_path, uni=True)
            pdf.set_font('SimHei', size=12)
        else:
            pdf.set_font('Arial', size=12)
        
        # 状态映射
        status_map = {
            'WAITING': '待接单', 'IN_PROGRESS': '维修中',
            'PENDING_ADMIN_CLOSE': '待审核', 'CLOSED': '已完成', 'REJECTED': '已取消'
        }
        
        # 添加标题
        pdf.add_page()
        pdf.set_font('SimHei', size=18)
        pdf.cell(0, 15, '报修工单报表', ln=True, align='C')
        pdf.ln(8)
        
        # 添加表头（横向A4宽度约297mm，减去左右边距约20mm，可用宽度约277mm）
        headers = ['报修编号', '设备名称', '设备类型', '位置', '报修人', '状态', '维修人员', '创建时间', '完成时间', '耗时']
        col_widths = [30, 36, 18, 48, 18, 16, 20, 28, 28, 16]
        
        pdf.set_font('SimHei', size=9)
        pdf.set_fill_color(220, 220, 220)
        for i, header in enumerate(headers):
            pdf.cell(col_widths[i], 7, header, border=1, fill=True, align='C')
        pdf.ln()
        
        # 添加数据行
        pdf.set_font('SimHei', size=8)
        for order in queryset[:100]:
            location_str = f"{order.location or ''}{order.building or ''}{order.floor or ''}{order.room or ''}"
            duration = ''
            if order.repair_start_time and order.complete_time:
                duration = str(round((order.complete_time - order.repair_start_time).total_seconds() / 60, 1))
            row_data = [
                (order.repair_no or '')[:12],
                (order.equipment_name[:18] if order.equipment_name else '')[:18],
                (str(order.equipment_type)[:10] if order.equipment_type else '')[:10],
                location_str[:25] if location_str else '',
                (order.user_name or '')[:8],
                status_map.get(order.status, order.status),
                (order.staff_name or '')[:8],
                order.created_at.strftime('%Y-%m-%d') if order.created_at else '',
                order.complete_time.strftime('%Y-%m-%d') if order.complete_time else '',
                duration
            ]
            for i, data in enumerate(row_data):
                pdf.cell(col_widths[i], 6, str(data), border=1, align='C')
            pdf.ln()
        
        # 保存到内存缓冲区
        buffer = io.BytesIO()
        buffer.write(pdf.output())
        buffer.seek(0)
        
        # 返回PDF文件
        response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
        response['Content-Disposition'] = 'attachment; filename="repair_orders.pdf"'
        return response
    
    @action(detail=False, methods=['get'])
    def spare_parts_excel(self, request):
        """
        导出备件库存Excel
        
        请求方式：
            GET /api/common/export/spare_parts_excel/
        
        返回：
            Excel文件下载
        """
        queryset = SparePart.objects.all()
        
        # 创建Excel工作簿
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '备件库存'
        
        # 设置表头
        headers = ['备件编号', '名称', '规格型号', '单位', '库存数量', '单价', '存放位置']
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)

        # 填充数据
        for row, part in enumerate(queryset, 2):
            ws.cell(row=row, column=1, value=part.part_no)
            ws.cell(row=row, column=2, value=part.name)
            ws.cell(row=row, column=3, value=part.specification)
            ws.cell(row=row, column=4, value=part.unit)
            ws.cell(row=row, column=5, value=part.quantity)
            ws.cell(row=row, column=6, value=float(part.unit_price))
            ws.cell(row=row, column=7, value=part.storage_location)
        
        # 保存到内存缓冲区
        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)
        
        # 返回Excel文件
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        response['Content-Disposition'] = 'attachment; filename="spare_parts.xlsx"'
        return response
